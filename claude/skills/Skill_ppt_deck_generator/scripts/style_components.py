#!/usr/bin/env python3
"""
style_components.py — reusable "signature" graphic components for reference style packs.

A reference's fidelity comes from COMPOSING these shared components (parametrised by its
palette + extracted assets) rather than writing bespoke code each time. build_pptx.py
calls them when a style pack ("fidelity mode") is active. Adding a new reference is then
mostly configuration in its style.json, not new drawing code.

Every function is palette-agnostic: pass hex colours / an asset path and it draws.
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SLIDE_W = 13.333
SLIDE_H = 7.5


def hex_to_rgb(h):
    h = (h or "#000000").lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _no_line_shadow(shp):
    shp.line.fill.background()
    shp.shadow.inherit = False


def _set_alpha(color, pct):
    """pct = opacity 0-100 (100 = opaque)."""
    from pptx.oxml.ns import qn
    srgb = color._xFill.find(qn('a:srgbClr'))
    if srgb is None:
        return
    srgb.append(srgb.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))}))


def fullbleed_image(slide, path):
    """Place an image covering the whole slide (add FIRST, before any text)."""
    slide.shapes.add_picture(path, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))


def scrim(slide, hexcolor, opacity_pct):
    """A translucent full-bleed overlay so text stays legible over imagery."""
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    r.fill.solid(); r.fill.fore_color.rgb = hex_to_rgb(hexcolor)
    _set_alpha(r.fill.fore_color, opacity_pct)
    _no_line_shadow(r)
    return r


def gradient(shape, c1, c2, angle_deg=90):
    """Apply a 2-stop linear gradient; fall back to solid c1 if unsupported."""
    try:
        shape.fill.gradient()
        stops = shape.fill.gradient_stops
        stops[0].color.rgb = hex_to_rgb(c1); stops[0].position = 0.0
        stops[1].color.rgb = hex_to_rgb(c2); stops[1].position = 1.0
        try:
            shape.fill.gradient_angle = angle_deg
        except Exception:
            pass
    except Exception:
        shape.fill.solid(); shape.fill.fore_color.rgb = hex_to_rgb(c1)


def rounded_card(slide, x, y, w, h, *, fill1, fill2=None, radius=0.08, line=None):
    """Rounded rectangle, optional vertical gradient (fill1→fill2) — Premium-style card."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if fill2:
        gradient(shp, fill1, fill2, 90)
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = hex_to_rgb(fill1)
    if line:
        shp.line.color.rgb = hex_to_rgb(line); shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def pill(slide, x, y, text, *, fill, textcolor, font, size=11, pad=0.16):
    """A rounded pill label (Premium's purple pill). Width auto-estimated from text."""
    w = max(0.9, pad * 2 + len(text) * size * 0.011)
    h = 0.42
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    shp.fill.solid(); shp.fill.fore_color.rgb = hex_to_rgb(fill)
    _no_line_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.name = font
    r.font.color.rgb = hex_to_rgb(textcolor)
    return w


def chevron(slide, x, y, hexcolor, font, size=20):
    """The Premium '>' mark, e.g. bottom-left of content slides."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(0.5), Inches(0.5))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run(); r.text = ">"
    r.font.size = Pt(size); r.font.bold = True; r.font.name = font
    r.font.color.rgb = hex_to_rgb(hexcolor)
    return tb


def _rt(slide, x, y, w, h, runs, *, font, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        line_spacing=1.12, letter=None):
    """runs: list of (text, size_pt, hex_color, bold). Simple multi-run textbox."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = line_spacing
    for (txt, size, col, bold) in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
        r.font.color.rgb = hex_to_rgb(col)
        if letter is not None:
            r.font._rPr.set('spc', str(int(letter * 100)))
    return tb


def numbered_eyebrow(slide, x, y, num, label, *, num_color, label_color, font):
    """Premium content eyebrow: big number + small label below (replaces the rule+label)."""
    if num:
        _rt(slide, x, y, 6, 0.6, [(num, 24, num_color, False)], font=font, letter=2.0)
    if label:
        _rt(slide, x, y + 0.5, 7.5, 0.35, [(label, 11, label_color, True)], font=font)


def sublabel(slide, x, y, w, text, *, color, font):
    """Purple UPPERCASE group label that heads a set of bullets (Premium right column)."""
    _rt(slide, x, y, w, 0.35, [((text or "").upper(), 12, color, True)], font=font, letter=1.2)


def bullet(slide, x, y, w, text, *, dot_color, text_color, font, size=13, marker="dot"):
    """A bullet row: a marker (filled dot or check ✓) + text. De-mix: the marker
    style is per-reference (Premium dot, Dark check)."""
    if marker == "check":
        _rt(slide, x - 0.02, y - 0.03, 0.34, 0.4, [("✓", size + 3, dot_color, True)], font=font)
    else:
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.09), Inches(0.09), Inches(0.09))
        d.fill.solid(); d.fill.fore_color.rgb = hex_to_rgb(dot_color); _no_line_shadow(d)
    _rt(slide, x + 0.3, y, w - 0.3, 0.7, [(text, size, text_color, False)],
        font=font, line_spacing=1.25)


def callout_box(slide, x, y, w, h, text, *, border, textcolor, font, fill="#FFFFFF"):
    """Rounded box with a purple border + left accent bar + purple text — Premium key phrase."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = 0.10
    except Exception:
        pass
    shp.fill.solid(); shp.fill.fore_color.rgb = hex_to_rgb(fill)
    shp.line.color.rgb = hex_to_rgb(border); shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.2), Inches(y + 0.22),
                                 Inches(0.06), Inches(h - 0.44))
    bar.fill.solid(); bar.fill.fore_color.rgb = hex_to_rgb(border); _no_line_shadow(bar)
    _rt(slide, x + 0.42, y, w - 0.62, h, [(text, 13, textcolor, True)],
        font=font, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.25)


def focus_tab(slide, x, y, text, *, fill, textcolor, font):
    """Small rotated purple 'Focus' tab overlapping a headline (Premium accent)."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(1.25), Inches(0.5))
    shp.fill.solid(); shp.fill.fore_color.rgb = hex_to_rgb(fill); _no_line_shadow(shp)
    shp.rotation = -10
    tf = shp.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(13); r.font.bold = True; r.font.name = font
    r.font.color.rgb = hex_to_rgb(textcolor)


def bento_card(slide, x, y, w, h, *, c1, c2):
    """Rounded purple-gradient bento card (Premium value/capability tiles)."""
    return rounded_card(slide, x, y, w, h, fill1=c1, fill2=c2, radius=0.07)


def make_donut(path, values, colors, *, size=900, hole=0.58):
    """Render an anti-aliased donut PNG (transparent hole). colors: hex strings."""
    from PIL import Image, ImageDraw
    ss = 3  # supersample for smooth edges
    S = size * ss
    im = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    total = float(sum(values)) or 1.0
    start = -90.0
    pad = int(S * 0.02)
    box = [pad, pad, S - pad, S - pad]
    for v, c in zip(values, colors):
        sweep = 360.0 * (v / total)
        d.pieslice(box, start, start + sweep, fill=c)
        start += sweep
    r = int(S * hole / 2); cx = S // 2
    d.ellipse([cx - r, cx - r, cx + r, cx + r], fill=(0, 0, 0, 0))
    im.resize((size, size), Image.LANCZOS).save(path)
